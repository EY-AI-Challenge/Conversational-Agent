from pathlib import Path
from typing import Iterable
import csv

from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_core.documents import Document
from openpyxl import load_workbook


PARTNER_NAME_HINTS = {
    "anabelasilva",
    "susanalencastre",
    "miguelcardosopinto",
    "fernandamarcelo",
    "davidgonçalves",
    "antoniooliveira",
    "antóniooliveira",
    "miguelamado",
    "ricardomorais",
    "jaimerocha",
    "ruimartins",
    "telmafranco",
    "manuelmota",
    "ananeves",
    "joãomoura",
    "joaomoura",
    "silviasilva",
    "rosáliaamorim",
    "rosaliaamorim",
    "ruihenriques",
    "nunocosta",
    "carlasápereira",
    "carlasapereira",
    "inêscabral",
    "inescabral",
    "saragomesferreira",
    "filipebrás",
    "filipebras",
    "isabelfaria",
    "joãosequeira",
    "joaosequeira",
    "ritacosta",
    "nunodeus",
    "jorgelibório",
    "jorgeliborio",
    "sérgioferreira",
    "sergioferreira",
    "joãonóbrega",
    "joaonobrega",
    "amilcarnunes",
    "luisflorindo",
    "luispedromendes",
    "teresafreitas",
    "inêsvazpereira",
    "inesvazpereira",
}


def classify_document(path: Path) -> str:
    name = path.stem.lower().replace(" ", "")
    full_name = path.name.lower()

    if "uploads" in path.parts:
        return "uploaded_document"
    if path.suffix.lower() in {".xlsx", ".xls", ".csv"}:
        return "service_line"
    if full_name.startswith("speech_to_text") or path.suffix.lower() == ".txt":
        return "transcript"
    if name in PARTNER_NAME_HINTS:
        return "partner_cv"
    if any(keyword in full_name for keyword in ["portugal", "ey.ai", "relatórios", "carreiras", "experience"]):
        return "news"
    return "internal"


def base_metadata(path: Path) -> dict:
    return {
        "source": path.name,
        "document_type": classify_document(path),
        "title": path.stem,
    }


def add_searchable_header(doc: Document) -> Document:
    source = doc.metadata.get("source", "")
    title = doc.metadata.get("title", "")
    document_type = doc.metadata.get("document_type", "")
    header = f"Documento: {source}\nTitulo: {title}\nTipo: {document_type}\n\n"
    if doc.page_content.startswith(header):
        return doc
    doc.page_content = f"{header}{doc.page_content}"
    return doc


def load_pdf(path: Path) -> list[Document]:
    docs = PyPDFLoader(str(path)).load()
    metadata = base_metadata(path)
    for doc in docs:
        doc.metadata.update(metadata)
        if "page" in doc.metadata and doc.metadata["page"] is not None:
            doc.metadata["page"] = int(doc.metadata["page"]) + 1
        add_searchable_header(doc)
    return docs


def load_text(path: Path) -> list[Document]:
    docs = TextLoader(str(path), encoding="utf-8", autodetect_encoding=True).load()
    metadata = base_metadata(path)
    for doc in docs:
        doc.metadata.update(metadata)
        add_searchable_header(doc)
    return docs


def _row_to_text(row: dict) -> str:
    parts = []
    for column, value in row.items():
        if value is not None and str(value).strip():
            parts.append(f"{column}: {str(value).strip()}")
    return "\n".join(parts)


def load_excel(path: Path) -> list[Document]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    documents: list[Document] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        headers = [str(cell).strip() if cell is not None else f"Column {index + 1}" for index, cell in enumerate(rows[0])]
        for row_index, values in enumerate(rows[1:], start=2):
            row = dict(zip(headers, values))
            text = _row_to_text(row)
            if not text:
                continue

            service_line = None
            for column, value in row.items():
                normalized_column = column.lower().replace("-", " ")
                if "service line" in normalized_column and value is not None:
                    service_line = str(value).strip()
                    break

            metadata = base_metadata(path)
            metadata.update(
                {
                    "sheet": sheet.title,
                    "row": int(row_index),
                    "service_line": service_line,
                }
            )
            documents.append(add_searchable_header(Document(page_content=text, metadata=metadata)))
    return documents


def load_csv(path: Path) -> list[Document]:
    documents: list[Document] = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row_index, row in enumerate(reader, start=2):
            text = _row_to_text(row)
            if not text:
                continue
            metadata = base_metadata(path)
            metadata.update({"row": row_index})
            documents.append(add_searchable_header(Document(page_content=text, metadata=metadata)))
    return documents


def load_document(path: Path) -> list[Document]:
    extension = path.suffix.lower()
    if extension == ".pdf":
        return load_pdf(path)
    if extension in {".txt", ".md"}:
        return load_text(path)
    if extension in {".xlsx", ".xls", ".csv"}:
        if extension == ".csv":
            return load_csv(path)
        return load_excel(path)
    return []


def iter_supported_files(data_dir: Path) -> Iterable[Path]:
    supported = {".pdf", ".txt", ".md", ".xlsx", ".xls", ".csv"}
    for path in sorted(data_dir.rglob("*")):
        if path.is_file() and path.suffix.lower() in supported:
            yield path
